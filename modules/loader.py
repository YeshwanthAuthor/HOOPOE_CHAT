import os
from langchain_core.documents import Document
from typing import List
import docx
import pypdf as pdf

from modules.utils import get_logger

logger = get_logger(__name__)


def text_to_doc(text: str, source: str, filename: str) -> Document:
    return Document(page_content=text, metadata={"source": source, "filename": filename})


def load_txt(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    logger.info("Loaded .txt file %s (%d characters)", os.path.basename(path), len(text))
    return [text_to_doc(text, os.path.abspath(path), os.path.basename(path))]


def load_docx(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    doc = docx.Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n\n".join(paragraphs)
    logger.info("Loaded .docx file %s (%d paragraphs)", os.path.basename(path), len(paragraphs))
    return [text_to_doc(text, os.path.abspath(path), os.path.basename(path))]


def load_pdf(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    pdf_data = pdf.PdfReader(path)
    docs = []
    filename = os.path.basename(path)
    for i, page in enumerate(pdf_data.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(text_to_doc(text, f"{os.path.abspath(path)}:page:{i + 1}", filename))
    logger.info("Loaded .pdf file %s (%d of %d pages had readable text)", filename, len(docs), len(pdf_data.pages))
    return docs


def load_document(path: str) -> list[Document] | None:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    allowed_extensions = [".txt", ".docx", ".pdf"]
    file_extension = os.path.splitext(path)[1].lower()
    if file_extension not in allowed_extensions:
        raise ValueError(
            f"File extension '{file_extension}' is not supported. "
            f"Allowed extensions are: {', '.join(allowed_extensions)}."
        )
    try:
        if file_extension == ".txt":
            return load_txt(path)
        elif file_extension == ".docx":
            return load_docx(path)
        elif file_extension == ".pdf":
            return load_pdf(path)
    except Exception as e:
        logger.error("Failed to load %s: %s", os.path.basename(path), e)
        raise RuntimeError(f"Failed to read {os.path.basename(path)}: {e}") from e
